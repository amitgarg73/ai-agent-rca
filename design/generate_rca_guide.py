"""
Generate the RCA View user guide as a .docx file.
Run: pip install python-docx && python3 design/generate_rca_guide.py
Output: design/RCA_View_Guide.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = os.path.join(os.path.dirname(__file__), "RCA_View_Guide.docx")

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1.2)


def h1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p


def h2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
    return p


def h3(text):
    p = doc.add_heading(text, level=3)
    p.runs[0].font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.runs[0].font.size = Pt(11)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(11)
        p.add_run(" — " + text).font.size = Pt(11)
    else:
        p.add_run(text).font.size = Pt(11)
    return p


def note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    run.italic = True
    return p


# ── Title page ────────────────────────────────────────────────────────────────
title_p = doc.add_heading("AI Agent RCA View", 0)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

sub = doc.add_paragraph("How to Read the Root Cause Analysis Dashboard")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(14)
sub.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

doc.add_paragraph("")
note_p = doc.add_paragraph("AI Agent Reliability — Observability Prototype")
note_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
note_p.runs[0].font.size = Pt(10)
note_p.runs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

doc.add_page_break()


# ── Overview ──────────────────────────────────────────────────────────────────
h1("Overview")
body(
    "The RCA View is the central diagnostic tool in the AI Agent Reliability dashboard. "
    "When an incident is detected, this page walks you from a plain-English summary of what "
    "went wrong, through the visual evidence, down to individual agent traces and exact fix steps."
)
body(
    "The page is built around one principle: show the answer first, then the evidence. "
    "Most engineers who arrive here want to know what broke and how to fix it in under "
    "30 seconds. The sections are ordered to support that."
)

doc.add_paragraph("")

h2("How to Navigate")
bullet("Open Incidents Feed from the sidebar.")
bullet("Click RCA on any incident row.")
bullet("The RCA View loads pre-populated with that incident.")
bullet("Use the incident selector at the top to switch between incidents.")
bullet("Use the share link (top of page) to send a direct URL to a teammate.")

doc.add_page_break()


# ── Section 1 ─────────────────────────────────────────────────────────────────
h1("Section 1: What Happened & What to Fix")

body(
    "This is the first thing you see. It gives you the complete story before you look at any "
    "trace evidence."
)

h2("Left panel — Incident summary")
bullet("Severity badge", "CRITICAL / WARNING / INFO")
bullet("The named failure pattern", "Tool Timeout Loop, Context Spiral, Pipeline Break, etc.")
bullet("A plain-English sentence stating the root cause, session duration, cost, and trade count.")
note("If trades = 0 and cost > $0, the entire session was waste.")

h2("Right panel — Fix steps")
body(
    "Numbered action items tailored to the failure pattern. These are not generic advice — "
    "they reference the specific tool, agent, and parameters that caused the failure."
)
bullet("Tool Timeout Loop: add timeout + retry with backoff, add fallback data source.")
bullet("Context Spiral: set token budget, add a decide-or-abort checkpoint.")
bullet("Pipeline Break: wrap the research-to-orchestrator handoff in try/except.")
bullet("Silent Exit: add terminal_reason logging to every orchestrator exit path.")
bullet("Cost Anomaly: find the high-cost agent, add per-session budget guard.")
bullet("Empty Result Loop: validate result quality after each tool call, fail fast.")

note(
    "Tip: read the fix steps first, then use the Agent Breakdown (Section 5) to confirm "
    "your understanding against the raw trace evidence."
)

doc.add_page_break()


# ── Section 2 ─────────────────────────────────────────────────────────────────
h1("Section 2: Incident Details")

body("Four KPI metrics and a cost breakdown.")

h2("KPI metrics")
bullet("Cost wasted", "dollars spent in this session with no useful output (0 trades)")
bullet("Tokens wasted", "total LLM tokens consumed across all agents in the session")
bullet("Trades", "number of trades the orchestrator ultimately executed")
bullet("Failed evals", "count of quality checks that did not meet their threshold")

h2("Session-level evals")
body(
    "These evals assess the pipeline as a whole, not individual agents. They run after the "
    "session completes."
)
bullet("pipeline_completion", "did all four required agents run? (market, research, risk, orchestrator)")
bullet("cost_anomaly", "was this session's cost more than 2 standard deviations above the rolling mean?")
bullet("outcome_linkage", "did the session produce a measurable outcome (trade or explicit terminal reason)?")
bullet("tokens_per_decision", "total tokens divided by trades executed — a waste efficiency metric")
note("Score 0.0 to 1.0. Green = passed threshold. Red = failed.")

h2("Cost by agent (donut chart)")
body(
    "Shows what fraction of the session's total cost each agent consumed. "
    "If one agent dominates the cost and the session wasted money, that agent is your primary "
    "optimisation target — even if it is not the root cause."
)

doc.add_page_break()


# ── Section 3 ─────────────────────────────────────────────────────────────────
h1("Section 3: Call Chain")

body(
    "A visual map of the agent pipeline showing at a glance where the failure occurred. "
    "Each circle is one agent. Arrows show the direction data flows (left to right)."
)

h2("Node colors")
bullet("Green", "all quality evals passed for this agent — healthy")
bullet("Red", "one or more evals failed — this agent has quality issues")
bullet("Amber", "root cause: this is where the failure originated")
bullet("Gray", "agent ran but has no eval data, or was skipped entirely")

h2("How to use it")
body(
    "Scan left to right. The first non-green node is where the pipeline started degrading. "
    "The amber node is where it definitively broke."
)
body(
    "Hover over any node to see a tooltip with token count, latency, and eval pass/fail count "
    "for that agent."
)
note(
    "A Pipeline Break incident will have a green market node, a green or amber research node, "
    "and a gray orchestrator node — the pipeline never reached the orchestrator."
)

doc.add_page_break()


# ── Section 4 ─────────────────────────────────────────────────────────────────
h1("Section 4: Execution Timeline")

body(
    "A horizontal Gantt chart showing every agent step positioned at its actual time "
    "within the session."
)

h2("Reading the chart")
bullet("X-axis", "seconds from session start (total session duration shown at a glance)")
bullet("Y-axis", "agent name (one row per agent)")
bullet("Bar length", "how long that step took (latency)")
bullet("Red bars", "steps that errored")
bullet("Colored bars", "successful steps, color-coded by agent — same colors as the Call Chain")
bullet("Gaps between agents", "handoff time or pipeline stall between agents")

h2("What to look for")
body(
    "A Tool Timeout Loop will show a dense cluster of red bars on the research row, "
    "all very long (30,000ms+), with no orchestrator row at all."
)
body(
    "A Pipeline Break will show normal research bars ending abruptly with no risk or "
    "orchestrator rows — the pipeline stalled mid-execution."
)
body(
    "A Context Spiral will show an unusually long research row with many steps, often "
    "with no corresponding orchestrator row."
)
note("Hover over any bar for step name, exact start time, duration, and outcome.")

doc.add_page_break()


# ── Section 5 ─────────────────────────────────────────────────────────────────
h1("Section 5: Agent Breakdown")

body(
    "The full evidence layer. One expandable card per agent. Cards with failures are "
    "expanded by default; healthy agents are collapsed to reduce noise."
)

h2("Card header")
body(
    "Each card header shows: agent name, ROOT CAUSE tag (if applicable), eval pass/fail counts, "
    "number of LLM calls, number of tool calls, total tokens, total latency, and error count."
)

h2("Inside each card")

h3("LLM Calls")
body(
    "One row per LLM call. Shows: call type, model name, tokens in and out, latency, outcome. "
    "If the LLM call errored, an error expander appears below it with the full error message."
)

h3("Tool Calls")
body(
    "One row per tool call. Shows: tool name, latency, outcome. "
    "Repeated identical errors are collapsed into a single summary row "
    "(e.g. 'get_stock_data x8 — error, repeated, collapsed') to reduce noise. "
    "Click the expander to see the full error text."
)

h3("Quality Evals")
body(
    "Automated quality scores for this agent. Score range: 0.0 to 1.0."
)
bullet("Green bar with + icon", "score met or exceeded the threshold — agent healthy on this dimension")
bullet("Red bar with x icon", "score fell below threshold — this dimension needs attention")
body(
    "Below each eval score is a one-line explanation of how the score was calculated "
    "(e.g. '3/8 tool calls succeeded (37%) — threshold >= 80%'). "
    "This tells you exactly what the eval measured and why it passed or failed."
)

h3("Root cause callout")
body(
    "If this agent is where the failure originated, an amber banner appears at the top of the card: "
    "'ROOT CAUSE — [root cause description]'."
)

h3("Recommended Fix")
body(
    "Shown only on the root cause agent. Same content as Section 1, but positioned here "
    "so you can read the fix in context of the traces that caused the failure."
)

doc.add_page_break()


# ── Section 6 ─────────────────────────────────────────────────────────────────
h1("Section 6: Compare to Healthy Session")

body(
    "Side-by-side comparison of the incident session vs. the most recent session "
    "with no incidents detected."
)

h2("Left column — this session")
body(
    "Shows cost, latency, token count, and trades for the incident session. "
    "Next to each metric is a percentage delta vs. the healthy baseline."
)
bullet("Red delta", "this session performed worse (higher cost, slower, fewer trades)")
bullet("Green delta", "this session performed better")

h2("Right column — healthy baseline")
body(
    "The most recent session with zero detected incidents. Used as the normal operating reference point."
)

h2("How to use it")
body(
    "A +400% cost delta with 0 trades compared to a healthy session that cost $0.02 and "
    "produced 2 trades is a clear signal of complete waste. Use this comparison to quantify "
    "the business impact of the incident before reporting it."
)
note(
    "If all recent sessions have incidents, the comparison will show a placeholder message. "
    "This itself is a signal — it means the system has not had a clean run recently."
)

doc.add_page_break()


# ── Controls ──────────────────────────────────────────────────────────────────
h1("Page Controls")

h2("Incident selector")
body(
    "The dropdown at the top of the page lists all detected incidents sorted by most recent. "
    "Selecting a different incident reloads all six sections for that incident."
)

h2("Share link")
body(
    "The URL is updated with a sid= parameter when you view an incident. Copy and share the URL "
    "to send a colleague directly to the same RCA view. Opening the URL pre-loads the incident."
)

h2("Re-run Analysis button")
body(
    "Re-runs all 13 evals and pattern detection for the current session and refreshes the page. "
    "Use this after changing eval logic or after a session's trace data is updated."
)

h2("Notes")
body(
    "A free-text field for adding your own observations. Saved to session state "
    "(survives page reloads within the same browser session). Use it to document "
    "findings, fix status, or context for teammates."
)

doc.add_page_break()


# ── Eval reference ────────────────────────────────────────────────────────────
h1("Eval Reference")

body("13 evals run per session. Organized by agent scope.")

h2("Market agent evals")
bullet("data_completeness", "fraction of market traces that succeeded (threshold >= 0.7)")
bullet("data_freshness", "time from session start to first market trace (threshold: within 30 min)")

h2("Research agent evals")
bullet("completion", "LLM ran AND at least one tool call succeeded = 1.0; LLM ran but all tools failed = 0.3 (threshold >= 0.7)")
bullet("token_efficiency", "tokens used vs. 30,000-token limit — score = 1 - used/(2*limit)")
bullet("tool_success_rate", "fraction of tool calls that succeeded (threshold >= 80%)")

h2("Risk agent evals")
bullet("assessment_complete", "risk agent traces succeeded = 1.0; no traces = 0.0 (threshold = 1.0)")
bullet("within_parameters", "fraction of risk traces with no errors (threshold >= 0.9)")

h2("Orchestrator agent evals")
bullet("decision_made", "trades executed or terminal reason logged = 1.0; no decision = 0.4 (threshold >= 0.7)")
bullet("consistency", "orchestrator pipeline ran without contradictions (threshold >= 0.8)")

h2("Session-level evals")
bullet("pipeline_completion", "all 4 agents ran (threshold = 1.0)")
bullet("cost_anomaly", "session cost within 2 standard deviations of rolling mean (threshold = not anomalous)")
bullet("outcome_linkage", "session has a measurable outcome: trades or terminal reason (threshold > 0)")
bullet("tokens_per_decision", "total tokens / trades executed; rewards efficient decisions")

doc.add_page_break()


# ── Failure patterns ──────────────────────────────────────────────────────────
h1("Failure Pattern Reference")

patterns = [
    (
        "Tool Timeout Loop",
        "CRITICAL",
        "The same tool errors 3 or more times in a row with no retry limit.",
        "research agent calling get_stock_data 8 times, all ReadTimeout, each 30,000ms.",
        "Add timeout + max_retries=2 with exponential backoff. Return a structured error if retries exhausted.",
    ),
    (
        "Context Spiral",
        "WARNING",
        "Research agent used more than 40,000 tokens and produced 0 trades.",
        "Research calls yfinance 20 times, accumulates context, never decides, exits silently.",
        "Set a hard token budget. Add a decide-or-abort checkpoint after each information-gathering step.",
    ),
    (
        "Pipeline Break",
        "CRITICAL",
        "Research agent ran but orchestrator never started.",
        "Uncaught exception in the research-to-orchestrator handoff. No terminal_reason logged.",
        "Wrap handoff in try/except. Log terminal_reason on any uncaught exception.",
    ),
    (
        "Empty Result Loop",
        "WARNING",
        "A tool succeeded 3 or more times but returned empty or unusable results, and 0 trades were produced.",
        "get_stock_data returns empty DataFrame. Research retries 5 times. Never decides.",
        "Validate result quality after each tool call. Fail fast on empty results rather than retrying.",
    ),
    (
        "Cost Anomaly",
        "WARNING",
        "Session cost exceeded the rolling mean by more than 2 standard deviations.",
        "Normal sessions cost $0.02. This session cost $1.98 — 99x the mean.",
        "Find the high-cost agent in the cost donut. Add per-session cost budgets.",
    ),
    (
        "Silent Exit",
        "INFO",
        "0 trades, no errors, no terminal reason logged.",
        "Orchestrator ran, found no opportunity, exited without recording why.",
        "Add terminal_reason logging to every orchestrator exit path (no_opportunity, risk_rejected, etc.).",
    ),
]

for name, severity, description, example, fix in patterns:
    h2(f"{name} ({severity})")
    bullet("What it is", description)
    bullet("Example", example)
    bullet("Fix", fix)
    doc.add_paragraph("")


# ── Save ──────────────────────────────────────────────────────────────────────
doc.save(OUT)
print(f"Saved: {OUT}")
